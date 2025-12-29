from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt


BLUE_DARK = RGBColor(11, 61, 145)  # #0B3D91
BLUE = RGBColor(30, 136, 229)  # #1E88E5
TEXT_DARK = RGBColor(12, 35, 74)
TEXT_MUTED = RGBColor(78, 106, 155)
BG = RGBColor(248, 251, 255)
BORDER = RGBColor(207, 224, 251)


def add_card(slide, x, y, w, h, title, big_left, big_right, note, accent=BLUE_DARK):
    card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = BORDER
    card.line.width = Pt(1.2)

    # accent bar
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, x, y, Cm(0.22), h)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    # title
    tx = slide.shapes.add_textbox(x + Cm(0.5), y + Cm(0.3), w - Cm(0.8), Cm(0.9))
    p = tx.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    # big numbers row
    row = slide.shapes.add_textbox(x + Cm(0.5), y + Cm(1.2), w - Cm(1), Cm(1.4))
    tf = row.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r1 = p.add_run()
    r1.text = big_left
    r1.font.size = Pt(22)
    r1.font.bold = True
    r1.font.color.rgb = TEXT_MUTED
    r2 = p.add_run()
    r2.text = "  →  "
    r2.font.size = Pt(18)
    r2.font.bold = True
    r2.font.color.rgb = BLUE
    r3 = p.add_run()
    r3.text = big_right
    r3.font.size = Pt(22)
    r3.font.bold = True
    r3.font.color.rgb = BLUE_DARK

    # note
    nt = slide.shapes.add_textbox(x + Cm(0.5), y + Cm(2.6), w - Cm(1), Cm(1.0))
    p = nt.text_frame.paragraphs[0]
    p.text = note
    p.font.size = Pt(11.5)
    p.font.color.rgb = TEXT_MUTED


def main() -> None:
    assets = Path("/workspace/assets")
    img_flow = assets / "flow_distributed_to_integrated.png"
    img_mm = assets / "multimodal_framework_map.png"

    out_dir = Path("/workspace/ppt")
    out_dir.mkdir(parents=True, exist_ok=True)
    out_pptx = out_dir / "work_highlight_full_chain_evaluation_one_slide.pptx"

    prs = Presentation()
    prs.slide_width = Cm(33.867)  # 13.333 in
    prs.slide_height = Cm(19.05)  # 7.5 in

    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # background
    bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()

    # top header bar (thin)
    top = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Cm(0.25))
    top.fill.solid()
    top.fill.fore_color.rgb = BLUE_DARK
    top.line.fill.background()

    # Title
    title = slide.shapes.add_textbox(Cm(1.2), Cm(0.8), Cm(20), Cm(1.3))
    tf = title.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "工作亮点：全链路评测体系构建"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = TEXT_DARK

    # Subtitle
    sub = slide.shapes.add_textbox(Cm(1.2), Cm(2.05), Cm(20), Cm(0.9))
    p = sub.text_frame.paragraphs[0]
    p.text = "从文本到多模态、从基准集成到榜单生产，形成可复用、可扩展、可服务化的评测能力"
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_MUTED

    # Left highlight box (editable text)
    left = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Cm(1.2), Cm(3.2), Cm(13.8), Cm(13.6))
    left.fill.solid()
    left.fill.fore_color.rgb = RGBColor(255, 255, 255)
    left.line.color.rgb = BORDER
    left.line.width = Pt(1.2)

    # Left section title
    lt = slide.shapes.add_textbox(Cm(1.8), Cm(3.6), Cm(12.8), Cm(0.9))
    p = lt.text_frame.paragraphs[0]
    p.text = "全链路交付内容"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = BLUE_DARK

    # bullet content
    tb = slide.shapes.add_textbox(Cm(1.8), Cm(4.6), Cm(12.8), Cm(11.8))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True

    bullets = [
        "高效集成：12项科学推理、4项数学、3项长文本、1项指令遵循评测基准（统一口径 + 标准流程）",
        "效率优化：科学推理评测周期 7天 压缩至 30小时；多类型榜单评测从 ≥4人协作 降至 单人完成",
        "多模态从0到1：搭建统一评测框架，开发13个多模态榜单并实现流程标准化",
        "能力沉淀：覆盖文本到多模态、基础到强化训练模型，为模型迭代提供核心数据支撑；为评测服务化/部署轻量化打基础",
    ]

    for i, t in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = t
        p.level = 0
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK
        p.space_after = Pt(6)

    # KPI cards (center top)
    add_card(
        slide,
        Cm(15.6),
        Cm(3.2),
        Cm(8.5),
        Cm(4.0),
        "科学推理评测周期",
        "7天",
        "30小时",
        "效率提升约 82%（统一口径 + 流程优化）",
        accent=BLUE_DARK,
    )
    add_card(
        slide,
        Cm(24.4),
        Cm(3.2),
        Cm(8.2),
        Cm(4.0),
        "多类型榜单协作",
        "≥4人",
        "1人",
        "人力成本降低约 75%（标准化生产）",
        accent=BLUE,
    )

    add_card(
        slide,
        Cm(15.6),
        Cm(7.5),
        Cm(17.0),
        Cm(3.7),
        "多模态评测体系",
        "0 → 1",
        "13个榜单",
        "统一框架适配，多任务模板化，结果可追溯，支撑服务化演进",
        accent=BLUE_DARK,
    )

    # Images (bottom-right area): side-by-side to fit one slide
    img_y = Cm(11.5)
    img_w = Cm(8.25)
    img_gap = Cm(0.5)
    img_x1 = Cm(15.6)
    img_x2 = img_x1 + img_w + img_gap

    # mini labels
    lab1 = slide.shapes.add_textbox(img_x1, img_y - Cm(0.55), img_w, Cm(0.5))
    p = lab1.text_frame.paragraphs[0]
    p.text = "插图 1：分散式 → 一体化流程"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.font.bold = True

    lab2 = slide.shapes.add_textbox(img_x2, img_y - Cm(0.55), img_w, Cm(0.5))
    p = lab2.text_frame.paragraphs[0]
    p.text = "插图 2：多模态统一框架能力版图"
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.font.bold = True

    slide.shapes.add_picture(str(img_flow), img_x1, img_y, width=img_w)
    slide.shapes.add_picture(str(img_mm), img_x2, img_y, width=img_w)

    # small caption
    cap = slide.shapes.add_textbox(Cm(15.6), Cm(18.45), Cm(17.0), Cm(0.55))
    p = cap.text_frame.paragraphs[0]
    p.text = "插图为生成素材（SVG源文件在 assets/ 中，可替换/二次编辑）"
    p.font.size = Pt(10.5)
    p.font.color.rgb = TEXT_MUTED
    p.alignment = PP_ALIGN.RIGHT

    prs.save(out_pptx)


if __name__ == "__main__":
    main()

